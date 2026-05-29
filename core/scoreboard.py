"""
Snooker Scoreboard - Main Scoreboard Window
"""

import json
import os
import sys
from datetime import datetime

from PyQt6.QtCore import (
    Qt, QTimer, QPropertyAnimation, QRect, QEasingCurve,
    QParallelAnimationGroup, QPoint, QPauseAnimation,
    QSequentialAnimationGroup, QSize, QRectF,
)
from PyQt6.QtGui import QFont, QColor, QPixmap, QImage, QPainter, QIcon
from PyQt6.QtSvg import QSvgRenderer
from PyQt6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton, QVBoxLayout, QHBoxLayout,
    QLineEdit, QFrame, QInputDialog, QMessageBox, QFileDialog,
    QSpacerItem, QSizePolicy, QGraphicsOpacityEffect, QDialog,
    QTableWidget, QTableWidgetItem, QGridLayout,
)

from config import UI, Game, Resources, Styles, get_resource_path
from utils import enable_win_blur, GlobalHotkey, HSpacerAnimator
from widgets import ClickableSvg, BadgeGhost, ScoreboardOverlay


class SnookerScoreboard(QWidget):
    """Main Snooker Scoreboard application window."""

    def __init__(self):
        super().__init__()
        # Pre-create commonly used fonts to avoid repeated creation
        self._init_fonts()
        self.init_game()

    def _init_fonts(self):
        """Initialize reusable font objects."""
        self.font_name = QFont("BBC Reith Sans", 14, QFont.Weight.Bold)
        self.font_name.setStyleStrategy(QFont.StyleStrategy.PreferAntialias)

        self.font_score = QFont("BBC Reith Sans", 15)
        self.font_score_bold = QFont("BBC Reith Sans", 15, QFont.Weight.Bold)
        self.font_score_normal = QFont("BBC Reith Sans", 15, QFont.Weight.Normal)

        self.font_break = QFont("BBC Reith Sans", 12)
        self.font_arrow = QFont("DaytonaPro", 30)
        self.font_ball_label = QFont("DaytonaPro", 16, QFont.Weight.Bold)

        self.font_badge_large = QFont("BBC Reith Sans", 11, QFont.Weight.Bold)
        self.font_badge_large.setStyleStrategy(QFont.StyleStrategy.PreferAntialias)

        self.font_badge_small = QFont("BBC Reith Sans", 9, QFont.Weight.Bold)
        self.font_badge_small.setStyleStrategy(QFont.StyleStrategy.PreferAntialias)

    # ==================== Game Initialization ====================

    def init_game(self):
        """Initialize or reset game state."""
        self._init_game_state()

        best_of = self._get_best_of_input()
        if best_of is None:
            return

        self.total_frames = best_of
        first_break = self._get_first_break_input()
        if first_break is None:
            return

        self.first_break_player = first_break
        self.current_frame_break_player = self.first_break_player

        self._setup_window()
        self.init_ui()
        self.update_scores()
        self.update_remaining_display()

    def _init_game_state(self):
        """Initialize all game state variables."""
        self.frame_high_break_p1 = 0
        self.frame_high_break_p2 = 0
        self.is100 = False
        self.switch_cooldown = False
        self.ball_cooldown = False
        self.in_black_ball_decider = False
        self.score_cooldowns = {}
        self.frame_scores_p1 = []
        self.frame_scores_p2 = []
        self.break_scores_p1 = []
        self.break_scores_p2 = []
        self.p1_break_slots = {}
        self.p2_break_slots = {}

        self.score_p1 = 0
        self.score_p2 = 0
        self.break_p1 = 0
        self.break_p2 = 0
        self.frame_score_p1 = 0
        self.frame_score_p2 = 0
        self.highest_break_p1 = 0
        self.highest_break_p2 = 0
        self.match_high_break_start_p1 = 0
        self.match_high_break_start_p2 = 0
        self.break_push_extra = 0
        self.total_push_extra = 0

        self.red_remaining = Game.INITIAL_REDS
        self.colors_remaining = {2: 1, 3: 1, 4: 1, 5: 1, 6: 1, 7: 1}
        self.action_stack = []
        self.just_potted_last_red = False
        self.buttons = []
        self.ball_buttons = []

        # Animation tracking
        self._active_animations = []

    def _get_best_of_input(self) -> int | None:
        """Get best-of frames input from user."""
        while True:
            best_of, ok = QInputDialog.getInt(
                self, "输入对局长度", "BO？(奇数且≤99)", min=1, max=Game.MAX_FRAMES
            )
            if not ok:
                self._exit_app()
                return None
            if best_of % 2 == 1:
                return best_of
            QMessageBox.warning(self, "非法输入", "请输入一个不超过99的奇数。")

    def _get_first_break_input(self) -> int | None:
        """Get first break player input from user."""
        first_items = ["左侧球员", "右侧球员"]
        who_first, ok = QInputDialog.getItem(
            self, "选择先开球一方", "谁先开球？", first_items, 0, False
        )
        if not ok:
            self._exit_app()
            return None
        return 1 if str(who_first).startswith("左") else 2

    def _exit_app(self):
        """Clean exit when user cancels initialization."""
        self.close()
        QTimer.singleShot(0, QApplication.quit)

    def _setup_window(self):
        """Configure main window properties."""
        self.setWindowTitle("斯诺克记分牌")
        self.setGeometry(100, 100, UI.MAIN_WINDOW_WIDTH, 300)

        self.current_break = self.total_frames
        self.current_player = self.current_frame_break_player

        screen = QApplication.primaryScreen().availableGeometry()
        target_width = UI.MAIN_WINDOW_WIDTH
        target_height = int(target_width / 4.295 * 3)
        self.setFixedSize(target_width, target_height)

        margin = 5
        x = screen.right() - self.width() - margin
        y = screen.top() + margin
        self.move(x, y)

    # ==================== UI Construction ====================

    def init_ui(self):
        """Build the user interface."""
        icon_path = get_resource_path(Resources.ICON)
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))

        main_layout = QVBoxLayout()
        main_layout.setSpacing(0)
        main_layout.setContentsMargins(0, 0, 0, 0)

        self.ball_remaining_labels = {}

        # Create flag widgets
        self._create_flag_widgets()

        # Create yellow bar (top bar with names and scores)
        yellow_bar = self._create_yellow_bar()

        # Create break bar (bottom bar with break info)
        break_bar = self._create_break_bar()

        # Create ball buttons
        ball_layout = self._create_ball_buttons()

        # Create foul buttons
        foul_layout = self._create_foul_buttons()

        # Create control buttons
        control_layout = self._create_control_buttons()

        # Assemble main layout
        main_layout.addSpacing(10)
        main_layout.addWidget(yellow_bar)
        main_layout.addWidget(break_bar)
        main_layout.addStretch(6)
        main_layout.addLayout(ball_layout)
        main_layout.addSpacing(20)
        main_layout.addLayout(control_layout)
        main_layout.addStretch(4)
        main_layout.addLayout(foul_layout)
        main_layout.addStretch(4)

        # Red ball minus button
        btn_red_minus = QPushButton("红球 -1")
        btn_red_minus.clicked.connect(self.remove_red_ball)
        main_layout.addWidget(btn_red_minus)
        self.buttons.append(btn_red_minus)

        # Manual score buttons
        manual_score_layout = self._create_manual_score_buttons()
        main_layout.addStretch(4)
        main_layout.addLayout(manual_score_layout)

        # Frame score control
        frame_score_layout = self._create_frame_score_buttons()
        main_layout.addLayout(frame_score_layout)

        self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)

        # Statistics and export buttons
        self.btn_stats = QPushButton("统计信息")
        self.btn_stats.clicked.connect(self.show_statistics)
        self.btn_stats.setFixedHeight(48)
        main_layout.addWidget(self.btn_stats)

        self.btn_export = QPushButton("导出Excel")
        self.btn_export.clicked.connect(self.export_to_excel)
        self.btn_export.setFixedHeight(48)
        self.buttons.append(self.btn_export)
        main_layout.addWidget(self.btn_export)

        # Set button heights
        for btn in self.buttons:
            btn.setFixedHeight(48)

        self.setLayout(main_layout)
        self.update_arrow()

        # Create overlay
        self.break_bar = break_bar
        self.overlay = ScoreboardOverlay(yellow_bar, break_bar, main_window=self)
        self.overlay.move(100, 100)
        self.overlay.show()

        # Century animation label
        self._create_century_label()

        self.overlay.setWindowOpacity(1.0)
        self.overlay_x = self.overlay.x()
        self.overlay_y = self.overlay.y()

        # Install global hotkey
        QTimer.singleShot(0, self._install_global_hotkey)
        QTimer.singleShot(0, lambda: enable_win_blur(self.overlay, acrylic=True, tint=(120, 120, 120, 1)))

    def _create_flag_widgets(self):
        """Create flag selection widgets."""
        self.flag_p1 = ClickableSvg()
        flag_p1_path = get_resource_path(Resources.FLAG_P1_DEFAULT)
        if os.path.exists(flag_p1_path):
            self.flag_p1.load(flag_p1_path)
        self.flag_p1.clicked.connect(lambda: self.change_flag(1))

        self.flag_p2 = ClickableSvg()
        flag_p2_path = get_resource_path(Resources.FLAG_P2_DEFAULT)
        if os.path.exists(flag_p2_path):
            self.flag_p2.load(flag_p2_path)
        self.flag_p2.clicked.connect(lambda: self.change_flag(2))

        self.flag_p1_path = flag_p1_path
        self.flag_p2_path = flag_p2_path

        self.flag_p1.setStyleSheet("border-radius: 4px; background: transparent;")
        self.flag_p2.setStyleSheet("border-radius: 4px; background: transparent;")

    def _create_yellow_bar(self) -> QFrame:
        """Create the top yellow bar with player names and scores."""
        yellow_bar = QFrame()
        yellow_bar.setObjectName("YellowBar")
        yellow_bar.setStyleSheet(f"""
            #YellowBar {{
                background-color: {Styles.YELLOW_BAR_BG};
                border-top-left-radius: 8px;
                border-top-right-radius: 8px;
            }}
        """)
        yellow_bar.setFixedHeight(UI.YELLOW_BAR_HEIGHT)

        yellow_layout = QHBoxLayout(yellow_bar)
        yellow_layout.setSpacing(0)
        yellow_layout.setContentsMargins(0, 0, 0, 0)

        # Player 1 name input
        self.name_input_p1 = QLineEdit("Player 1")
        self.name_input_p1.setFont(self.font_name)
        self.name_input_p1.setStyleSheet(
            "border: none; background-color: transparent; color: black; padding-top: 6px;"
        )

        # Arrows
        self.arrow_left = QLabel("◀")
        self.arrow_left.setStyleSheet("background: transparent; padding-bottom: 2px;")
        self.arrow_left.setFont(self.font_arrow)
        self.arrow_left.setFixedWidth(30)
        self.arrow_left.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.arrow_right = QLabel("▶")
        self.arrow_right.setStyleSheet("background: transparent; padding-bottom: 2px;")
        self.arrow_right.setFont(self.font_arrow)
        self.arrow_right.setFixedWidth(30)
        self.arrow_right.setAlignment(Qt.AlignmentFlag.AlignCenter)

        # Score labels
        score_style = "background-color: white; border-left: 2px solid black; border-right: 3px solid black; padding-top: 6px;"
        self.label_score_p1 = QLabel("0")
        self.label_score_p1.setFont(self.font_score)
        self.label_score_p1.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label_score_p1.setFixedSize(75, UI.YELLOW_BAR_HEIGHT)
        self.label_score_p1.setStyleSheet(score_style)

        self.label_score_p2 = QLabel("0")
        self.label_score_p2.setFont(self.font_score)
        self.label_score_p2.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label_score_p2.setFixedSize(75, UI.YELLOW_BAR_HEIGHT)
        self.label_score_p2.setStyleSheet(
            "background-color: white; border-left: 3px solid black; border-right: 2px solid black; padding-top: 6px"
        )

        # Frame score labels
        frame_style = f"background-color: {Styles.TEAL_BG}; color: white; padding-top: 6px;"

        self.label_frame_left = QLabel(f"{self.frame_score_p1:>2}")
        self.label_frame_left.setFont(self.font_score)
        self.label_frame_left.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label_frame_left.setFixedSize(45, UI.YELLOW_BAR_HEIGHT)
        self.label_frame_left.setStyleSheet(frame_style)

        self.label_frame_center = QLabel(f"({self.total_frames})")
        self.label_frame_center.setFont(self.font_score)
        self.label_frame_center.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label_frame_center.setFixedSize(42, UI.YELLOW_BAR_HEIGHT)
        self.label_frame_center.setStyleSheet(frame_style)

        self.label_frame_right = QLabel(f"{self.frame_score_p2:>2}")
        self.label_frame_right.setFont(self.font_score)
        self.label_frame_right.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.label_frame_right.setFixedSize(45, UI.YELLOW_BAR_HEIGHT)
        self.label_frame_right.setStyleSheet(frame_style)

        # Player 2 name input
        self.name_input_p2 = QLineEdit("Player 2")
        self.name_input_p2.setFont(self.font_name)
        self.name_input_p2.setAlignment(Qt.AlignmentFlag.AlignRight)
        self.name_input_p2.setStyleSheet(
            "border: none; background-color: transparent; color: black; padding-top: 6px"
        )

        self.arrow_left_initial_pos = self.arrow_left.pos()
        self.arrow_right_initial_pos = self.arrow_right.pos()

        # Left side layout
        left_layout = QHBoxLayout()
        left_layout.setSpacing(0)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.addWidget(self.name_input_p1)
        left_layout.addSpacing(2)
        left_layout.addWidget(self.arrow_left)
        left_layout.addWidget(self.label_score_p1)

        # Right side layout
        right_layout = QHBoxLayout()
        right_layout.setSpacing(0)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.addWidget(self.label_score_p2)
        right_layout.addWidget(self.arrow_right)
        right_layout.addSpacing(2)
        right_layout.addWidget(self.name_input_p2)

        # Assemble yellow bar
        flag_p1_wrapper = QHBoxLayout()
        flag_p1_wrapper.setContentsMargins(8, 2, 0, 0)
        flag_p1_wrapper.addWidget(self.flag_p1)

        flag_p2_wrapper = QHBoxLayout()
        flag_p2_wrapper.setContentsMargins(0, 2, 8, 0)
        flag_p2_wrapper.addWidget(self.flag_p2)

        yellow_layout.addLayout(flag_p1_wrapper)
        yellow_layout.addSpacing(6)
        yellow_layout.addLayout(left_layout)
        yellow_layout.addWidget(self.label_frame_left)
        yellow_layout.addWidget(self.label_frame_center)
        yellow_layout.addWidget(self.label_frame_right)
        yellow_layout.addLayout(right_layout)
        yellow_layout.addSpacing(6)
        yellow_layout.addLayout(flag_p2_wrapper)

        return yellow_bar

    def _create_break_bar(self) -> QFrame:
        """Create the bottom break bar with break info and ball stats."""
        break_bar = QFrame()
        break_bar.setObjectName("BreakBar")
        break_bar.setStyleSheet(f"""
            #BreakBar {{
                background-color: {Styles.BREAK_BAR_BG};
                border-bottom-left-radius: 8px;
                border-bottom-right-radius: 8px;
            }}
        """)
        break_bar.setFixedHeight(UI.BREAK_BAR_HEIGHT)

        break_layout = QGridLayout(break_bar)
        break_layout.setContentsMargins(10, 0, 10, 0)

        # Left break info
        left_break = QVBoxLayout()
        left_break.setContentsMargins(0, 4, 0, 8)

        # Right break info
        right_break = QVBoxLayout()
        right_break.setContentsMargins(0, 4, 0, 8)

        break_layout.addLayout(left_break, 0, 0, alignment=Qt.AlignmentFlag.AlignLeft)

        # Middle info widget
        self.middle_info_widget = QWidget()
        self.middle_info_widget.setStyleSheet("background-color: rgba(68, 68, 68, 0)")
        middle_info_layout = QVBoxLayout(self.middle_info_widget)
        middle_info_layout.setContentsMargins(0, 0, 0, 0)
        middle_info_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.label_ahead = QLabel("Ahead: 0")
        self.label_remaining = QLabel("Remaining: 147")
        for lbl in [self.label_ahead, self.label_remaining]:
            lbl.setFont(self.font_break)
            lbl.setStyleSheet("color: white; background-color: rgba(68, 68, 68, 0);")
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            middle_info_layout.addWidget(lbl)

        break_layout.addWidget(self.middle_info_widget, 0, 1, alignment=Qt.AlignmentFlag.AlignCenter)
        break_layout.addLayout(right_break, 0, 2, alignment=Qt.AlignmentFlag.AlignRight)
        break_layout.setColumnStretch(0, 1)
        break_layout.setColumnStretch(1, 0)
        break_layout.setColumnStretch(2, 1)

        # Break and highest break labels
        self.label_break_p1 = QLabel("Break: 0")
        self.label_break_p1.setFont(self.font_break)
        self.label_break_p1.setStyleSheet("color: white; background-color: rgba(68, 68, 68, 0);")
        self.label_break_p1.setFixedWidth(100)

        self.label_high_break_p1 = QLabel("Highest: 0")
        self.label_high_break_p1.setFont(self.font_break)
        self.label_high_break_p1.setStyleSheet("color: white; background-color: rgba(68, 68, 68, 0);")
        self.label_high_break_p1.setFixedWidth(100)
        self.label_high_break_p1.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignLeft)

        self.label_break_p2 = QLabel("Break: 0")
        self.label_break_p2.setFont(self.font_break)
        self.label_break_p2.setStyleSheet("color: white; background-color: rgba(68, 68, 68, 0);")
        self.label_break_p2.setFixedWidth(100)

        self.label_high_break_p2 = QLabel("Highest: 0")
        self.label_high_break_p2.setFont(self.font_break)
        self.label_high_break_p2.setStyleSheet("color: white; background-color: rgba(68, 68, 68, 0);")
        self.label_high_break_p2.setFixedWidth(100)
        self.label_high_break_p2.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignRight)

        # Ball stats
        self._create_ball_stats()

        # Assemble left break area
        left_break_row = QHBoxLayout()
        left_break_row.addWidget(self.label_break_p1)
        left_break_row.addSpacing(20)
        left_break_row.addWidget(self.label_high_break_p1)
        left_break_row.addStretch()
        left_break.addLayout(left_break_row)

        p1_stat_wrapper = QHBoxLayout()
        p1_stat_wrapper.addLayout(self.p1_stat_layout)
        p1_stat_wrapper.addStretch()
        left_break.addLayout(p1_stat_wrapper)
        left_break.addSpacing(6)

        p1_total_wrapper = QHBoxLayout()
        p1_total_wrapper.addLayout(self.p1_total_stat_layout)
        p1_total_wrapper.addStretch()
        left_break.addLayout(p1_total_wrapper)

        # Assemble right break area
        right_break_row = QHBoxLayout()
        right_break_row.addStretch()
        right_break_row.addWidget(self.label_high_break_p2)
        right_break_row.addSpacing(20)
        right_break_row.addWidget(self.label_break_p2)
        right_break.addLayout(right_break_row)

        p2_stat_wrapper = QHBoxLayout()
        p2_stat_wrapper.addStretch()
        p2_stat_wrapper.addLayout(self.p2_stat_layout)
        right_break.addLayout(p2_stat_wrapper)
        right_break.addSpacing(6)

        p2_total_wrapper = QHBoxLayout()
        p2_total_wrapper.addStretch()
        p2_total_wrapper.addLayout(self.p2_total_stat_layout)
        right_break.addLayout(p2_total_wrapper)

        return break_bar

    def _create_ball_stats(self):
        """Create ball statistics labels for both players."""
        self.p1_ball_stats = {}
        self.p2_ball_stats = {}
        self.p1_stat_layout = QHBoxLayout()
        self.p2_stat_layout = QHBoxLayout()
        self.p1_frame_total_stats = {}
        self.p2_frame_total_stats = {}
        self.p1_total_stat_layout = QHBoxLayout()
        self.p2_total_stat_layout = QHBoxLayout()

        # Current break badges (large)
        for name, color, val in Game.BALL_CONFIG:
            for stats, layout in [
                (self.p1_ball_stats, self.p1_stat_layout),
                (self.p2_ball_stats, self.p2_stat_layout),
            ]:
                label = QLabel()
                label.setFont(self.font_badge_large)
                label.setFixedSize(UI.BADGE_SIZE_LARGE, UI.BADGE_SIZE_LARGE)
                label.setAlignment(Qt.AlignmentFlag.AlignCenter)
                text_color = "black" if val in Game.LIGHT_BALLS else "white"
                label.setStyleSheet(f"""
                    border-radius: {UI.BADGE_SIZE_LARGE // 2}px;
                    background-color: {color};
                    color: {text_color};
                    padding-top: 4px;
                """)
                layout.setSpacing(6)
                layout.setContentsMargins(6, 4, 6, 4)
                label.hide()
                layout.addWidget(label)
                stats[val] = label

        # Frame total badges (small)
        for name, color, val in Game.BALL_CONFIG:
            for stats, layout in [
                (self.p1_frame_total_stats, self.p1_total_stat_layout),
                (self.p2_frame_total_stats, self.p2_total_stat_layout),
            ]:
                label = QLabel()
                label.setFont(self.font_badge_small)
                label.setFixedSize(UI.BADGE_SIZE_SMALL, UI.BADGE_SIZE_SMALL)
                label.setAlignment(Qt.AlignmentFlag.AlignCenter)
                text_color = "black" if val in Game.LIGHT_BALLS else "white"
                label.setStyleSheet(
                    f"border-radius: {UI.BADGE_SIZE_SMALL // 2}px; background-color: {color}; color: {text_color}; padding-top: 2px;"
                )
                layout.setSpacing(4)
                layout.setContentsMargins(6, 0, 6, 0)
                label.hide()
                layout.addWidget(label)
                stats[val] = label

        self.p1_stat_layout.setSpacing(5)
        self.p2_stat_layout.setSpacing(5)
        self.p1_stat_layout.setContentsMargins(5, 0, 0, 0)
        self.p2_stat_layout.setContentsMargins(0, 0, 5, 0)

    def _create_ball_buttons(self) -> QHBoxLayout:
        """Create ball potting buttons."""
        ball_layout = QHBoxLayout()

        for name, color, value in Game.BALL_CONFIG:
            container = QFrame()
            container.setFixedSize(UI.BALL_BUTTON_SIZE, UI.BALL_BUTTON_SIZE)
            grid = QGridLayout(container)
            grid.setContentsMargins(0, 0, 0, 0)
            grid.setSpacing(0)

            btn = QPushButton()
            btn.setFixedSize(UI.BALL_BUTTON_SIZE, UI.BALL_BUTTON_SIZE)
            btn.setStyleSheet(f"border-radius: {UI.BALL_BUTTON_SIZE // 2}px; background-color: {color};")
            btn.clicked.connect(lambda checked, v=value, b=btn: self.handle_score_click(v, b))
            btn.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            self.ball_buttons.append(btn)
            self.score_cooldowns[value] = False

            label = QLabel("15" if value == 1 else "1")
            label.setAlignment(Qt.AlignmentFlag.AlignCenter)
            label.setFont(self.font_ball_label)
            text_color = "black" if value in Game.LIGHT_BALLS else "white"
            label.setStyleSheet(f"color: {text_color}; background: transparent;")
            label.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)

            grid.addWidget(btn, 0, 0)
            grid.addWidget(label, 0, 0)
            self.ball_remaining_labels[value] = label
            ball_layout.addWidget(container)

        return ball_layout

    def _create_foul_buttons(self) -> QHBoxLayout:
        """Create foul buttons."""
        foul_layout = QHBoxLayout()
        for value in [4, 5, 6, 7]:
            btn = QPushButton(f"罚{value}")
            btn.clicked.connect(lambda checked, v=value: self.add_foul(v))
            foul_layout.addWidget(btn)
            self.buttons.append(btn)
        return foul_layout

    def _create_control_buttons(self) -> QHBoxLayout:
        """Create control buttons."""
        control_layout = QHBoxLayout()

        self.btn_reset = QPushButton("重置")
        self.btn_reset.clicked.connect(self.reset_scores)
        self.buttons.append(self.btn_reset)

        self.btn_rerack = QPushButton("重摆")
        self.btn_rerack.clicked.connect(self.rerack_frame)
        self.buttons.append(self.btn_rerack)

        self.btn_switch = QPushButton("切换击球方")
        self.btn_switch.clicked.connect(self.switch_player)
        self.buttons.append(self.btn_switch)

        self.btn_undo = QPushButton("撤销")
        self.btn_undo.clicked.connect(self.undo_action)
        self.buttons.append(self.btn_undo)

        self.btn_full_reset = QPushButton("全部重置")
        self.btn_full_reset.clicked.connect(self.full_reset)
        self.buttons.append(self.btn_full_reset)

        self.btn_save = QPushButton("存档")
        self.btn_save.clicked.connect(self.save_game)
        self.buttons.append(self.btn_save)

        self.btn_load = QPushButton("加载")
        self.btn_load.clicked.connect(self.load_game)
        self.buttons.append(self.btn_load)

        control_layout.addWidget(self.btn_reset)
        control_layout.addWidget(self.btn_rerack)
        control_layout.addWidget(self.btn_switch)
        control_layout.addWidget(self.btn_undo)
        control_layout.addWidget(self.btn_full_reset)
        control_layout.addWidget(self.btn_save)
        control_layout.addWidget(self.btn_load)

        return control_layout

    def _create_manual_score_buttons(self) -> QHBoxLayout:
        """Create manual score buttons."""
        layout = QHBoxLayout()
        for value in range(1, 8):
            btn = QPushButton(f"+{value}")
            btn.clicked.connect(lambda checked, v=value: self.add_manual_score(v))
            layout.addWidget(btn)
            self.buttons.append(btn)
            self.score_cooldowns[value] = False
        return layout

    def _create_frame_score_buttons(self) -> QHBoxLayout:
        """Create frame score control buttons."""
        layout = QHBoxLayout()

        btn_frame_plus = QPushButton("对局比分 +1")
        btn_frame_plus.clicked.connect(self.increase_frame_score)
        layout.addWidget(btn_frame_plus)
        self.buttons.append(btn_frame_plus)

        btn_frame_minus = QPushButton("对局比分 -1")
        btn_frame_minus.clicked.connect(self.decrease_frame_score)
        layout.addWidget(btn_frame_minus)
        self.buttons.append(btn_frame_minus)

        return layout

    def _create_century_label(self):
        """Create the century celebration label."""
        self.century_label = QLabel("C E N T U R Y", self.break_bar)
        self.century_label.setFont(QFont("BBC Reith Sans"))
        self.century_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.century_label.setStyleSheet("""
            background-color: rgba(255, 255, 255, 255);
            color: black;
            font-size: 24px;
            font-weight: bold;
            border-radius: 10px;
            padding-top: 6px
        """)
        self.century_label.setGeometry(0, 0, 0, 0)
        self.century_label.hide()

    # ==================== Score Management ====================

    def increase_frame_score(self):
        """Increase current player's frame score."""
        if self.current_player == 1:
            self.frame_score_p1 += 1
        else:
            self.frame_score_p2 += 1
        self.update_scores()

    def decrease_frame_score(self):
        """Decrease current player's frame score."""
        if self.current_player == 1 and self.frame_score_p1 > 0:
            self.frame_score_p1 -= 1
        elif self.current_player == 2 and self.frame_score_p2 > 0:
            self.frame_score_p2 -= 1
        self.update_scores()

    def get_frame_score_text_left(self) -> str:
        return f"{self.frame_score_p1:>2}"

    def get_frame_score_text_right(self) -> str:
        return f"{self.frame_score_p2:<2}"

    def update_arrow(self):
        """Update arrow visibility based on current player."""
        self.arrow_left.setVisible(self.current_player == 1)
        self.arrow_right.setVisible(self.current_player == 2)

    def _compute_next_break_player(self) -> int:
        """Compute who should break in the next frame."""
        finished_frames = len(self.frame_scores_p1)
        if finished_frames % 2 == 0:
            return self.first_break_player
        return 2 if self.first_break_player == 1 else 1

    def update_remaining_display(self):
        """Update remaining points and ahead display."""
        if self.red_remaining > 0:
            points = self.red_remaining * 8 + sum(k * v for k, v in self.colors_remaining.items())
        else:
            points = sum(k * v for k, v in self.colors_remaining.items())

        diff = self.score_p1 - self.score_p2
        lead = abs(diff)
        ahead_text = f"Ahead: {lead}" if diff != 0 else "Ahead: 0"

        self.label_ahead.setText(ahead_text)
        self.label_remaining.setText(f"Remaining: {points}")

        is_over = lead > points
        color = Styles.ALERT_COLOR if is_over else "white"
        style = f"color: {color}; background-color: rgba(68, 68, 68, 0);"
        self.label_ahead.setStyleSheet(style)
        self.label_remaining.setStyleSheet(style)

    def update_scores(self):
        """Update all score displays."""
        self.label_score_p1.setText(str(self.score_p1))
        self.label_score_p2.setText(str(self.score_p2))
        self.label_frame_left.setText(self.get_frame_score_text_left())
        self.label_frame_right.setText(self.get_frame_score_text_right())

        # Update frame score font weight
        if self.frame_score_p1 > self.frame_score_p2:
            self.label_frame_left.setFont(self.font_score_bold)
            self.label_frame_right.setFont(self.font_score_normal)
        elif self.frame_score_p2 > self.frame_score_p1:
            self.label_frame_left.setFont(self.font_score_normal)
            self.label_frame_right.setFont(self.font_score_bold)
        else:
            self.label_frame_left.setFont(self.font_score_normal)
            self.label_frame_right.setFont(self.font_score_normal)

        self.label_break_p1.setText(f"Break: {self.break_p1}")
        self.label_break_p1.setAlignment(Qt.AlignmentFlag.AlignLeft)
        self.label_break_p2.setText(f"Break: {self.break_p2}")
        self.label_break_p2.setAlignment(Qt.AlignmentFlag.AlignRight)

        # Update highest break
        if self.current_player == 1:
            if self.break_p1 > self.frame_high_break_p1:
                self.frame_high_break_p1 = self.break_p1
            if self.break_p1 > self.highest_break_p1:
                self.highest_break_p1 = self.break_p1
            self.label_high_break_p1.setText(f"Highest: {self.highest_break_p1}")
        else:
            if self.break_p2 > self.frame_high_break_p2:
                self.frame_high_break_p2 = self.break_p2
            if self.break_p2 > self.highest_break_p2:
                self.highest_break_p2 = self.break_p2
            self.label_high_break_p2.setText(f"Highest: {self.highest_break_p2}")

    # ==================== Ball Stats ====================

    def update_ball_stats(self, value):
        """Update current break ball statistics."""
        stats = self.p1_ball_stats if self.current_player == 1 else self.p2_ball_stats
        label = stats.get(value)
        if not label:
            return

        count = label.text().strip()
        if not count:
            label.setText("1")
            label.hide()
            layout = self.p1_stat_layout if self.current_player == 1 else self.p2_stat_layout
            self._push_neighbors_then_fly_in(layout, label, align_right=(self.current_player == 2))
        else:
            new_text = str(int(count) + 1)
            self.animate_stat_update(label, new_text)
            label.show()

    def update_total_ball_stats(self, player, value):
        """Update frame total ball statistics."""
        stats = self.p1_frame_total_stats if player == 1 else self.p2_frame_total_stats
        layout = self.p1_total_stat_layout if player == 1 else self.p2_total_stat_layout
        label = stats.get(value)
        if not label:
            return

        txt = (label.text() or "").strip()
        if not txt:
            label.setText("1")
            label.hide()
            idx = self._layout_index_of(layout, label)
            if idx < 0:
                self.animate_appearance(label, from_below=True)
                return

            align_right = player == 2
            if not self._need_push(layout, label, align_right):
                self.animate_appearance(label, from_below=True)
                return

            insert_at = idx + 1 if align_right else idx
            h = label.height() or label.sizeHint().height() or 20
            spacer = QSpacerItem(0, h, QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
            layout.insertSpacerItem(insert_at, spacer)

            badge_w = label.width() or label.sizeHint().width() or 20
            gap = layout.spacing()
            extra = getattr(self, "total_push_extra", 0)
            target_w = badge_w + gap + extra

            anim_obj = HSpacerAnimator(layout, spacer, h, parent=self)
            anim = QPropertyAnimation(anim_obj, b"width", self)
            anim.setDuration(200)
            anim.setStartValue(0)
            anim.setEndValue(max(0, target_w))
            anim.setEasingCurve(QEasingCurve.Type.OutCubic)

            def _finish():
                for i in range(layout.count()):
                    it = layout.itemAt(i)
                    if it and it.spacerItem() is spacer:
                        layout.takeAt(i)
                        break
                self.animate_appearance(label, from_below=True)

            anim.finished.connect(_finish)
            self._track_animation(anim)
            anim.start()
        else:
            self.animate_stat_update(label, str(int(txt) + 1), scale=1.12)
            label.show()

    def decrement_total_ball_stats(self, player, value):
        """Decrement frame total ball statistics."""
        stats = self.p1_frame_total_stats if player == 1 else self.p2_frame_total_stats
        label = stats.get(value)
        if not label:
            return

        count = label.text()
        if count and int(count) > 1:
            label.setText(str(int(count) - 1))
        else:
            label.setText("")
            label.hide()

    def clear_break_stats_for_player(self, player):
        """Clear break statistics for a player."""
        stats = self.p1_ball_stats if player == 1 else self.p2_ball_stats
        for label in stats.values():
            label.setText("")
            label.hide()

    def clear_total_stats_both(self):
        """Clear total statistics for both players."""
        for stats in [self.p1_frame_total_stats, self.p2_frame_total_stats]:
            for label in stats.values():
                label.setText("")
                label.hide()

    def snapshot_break_stats(self) -> tuple:
        """Take a snapshot of current break statistics."""
        snap_p1 = {v: int(lbl.text()) if lbl.text() else 0 for v, lbl in self.p1_ball_stats.items()}
        snap_p2 = {v: int(lbl.text()) if lbl.text() else 0 for v, lbl in self.p2_ball_stats.items()}
        return snap_p1, snap_p2

    def restore_break_stats_from_snapshot(self, snap_p1, snap_p2):
        """Restore break statistics from a snapshot."""
        for v, lbl in self.p1_ball_stats.items():
            n = snap_p1.get(v, 0)
            if n > 0:
                lbl.setText(str(n))
                lbl.show()
            else:
                lbl.setText("")
                lbl.hide()
        for v, lbl in self.p2_ball_stats.items():
            n = snap_p2.get(v, 0)
            if n > 0:
                lbl.setText(str(n))
                lbl.show()
            else:
                lbl.setText("")
                lbl.hide()

    # ==================== Ball Potting ====================

    def pot_ball(self, value):
        """Handle potting a ball."""
        if self.ball_cooldown:
            return

        self.ball_cooldown = True
        QTimer.singleShot(UI.BALL_COOLDOWN, lambda: setattr(self, "ball_cooldown", False))

        if self.in_black_ball_decider and value != 7:
            return

        # Check if ball can be potted
        if value == 1 and self.red_remaining == 0:
            return
        if (value != 1 and self.red_remaining == 0 and not self.just_potted_last_red
                and value in self.colors_remaining and self.colors_remaining[value] == 0):
            return

        # Save action for undo
        self.action_stack.append((
            "pot", value, self.current_player, self.red_remaining,
            dict(self.colors_remaining), self.just_potted_last_red,
            self.highest_break_p1, self.highest_break_p2, self.is100,
        ))

        # Add score
        if self.current_player == 1:
            self.score_p1 += value
            self.break_p1 += value
        else:
            self.score_p2 += value
            self.break_p2 += value

        # Update ball remaining
        if value == 1:
            if self.red_remaining > 0:
                self.red_remaining -= 1
                if self.red_remaining == 0:
                    self.just_potted_last_red = True
            self.ball_remaining_labels[1].setText(str(self.red_remaining))
        elif self.red_remaining == 0 and value in self.colors_remaining:
            if self.just_potted_last_red:
                self.just_potted_last_red = False
            elif value in self.colors_remaining and self.colors_remaining[value] > 0:
                self.colors_remaining[value] -= 1
            self.ball_remaining_labels[value].setText(str(self.colors_remaining[value]))

        self.update_scores()
        self.update_ball_stats(value)
        self.update_total_ball_stats(self.current_player, value)
        self.update_remaining_display()

        # Check for black ball decider
        if (not self.in_black_ball_decider and self.red_remaining == 0
                and all(v == 0 for v in self.colors_remaining.values())):
            if self.score_p1 == self.score_p2:
                self.in_black_ball_decider = True
                self.colors_remaining[7] = 1
                QMessageBox.information(self, "争黑球", "双方分数相同，进入争黑球阶段！只能击打黑球。")

        # Check for century
        current_break = self.break_p1 if self.current_player == 1 else self.break_p2
        if current_break >= Game.CENTURY_BREAK and not self.is100:
            self.show_century_animation()
            self.is100 = True

    def add_manual_score(self, value):
        """Add manual score without affecting ball count."""
        if self.ball_cooldown:
            return

        self.ball_cooldown = True
        QTimer.singleShot(UI.MANUAL_SCORE_COOLDOWN, lambda: setattr(self, "ball_cooldown", False))

        self.action_stack.append((
            "manual", value, self.current_player,
            self.highest_break_p1, self.highest_break_p2,
        ))

        if self.current_player == 1:
            self.score_p1 += value
            self.break_p1 += value
        else:
            self.score_p2 += value
            self.break_p2 += value

        self.update_scores()
        self.update_ball_stats(value)
        self.update_total_ball_stats(self.current_player, value)

        current_break = self.break_p1 if self.current_player == 1 else self.break_p2
        if current_break >= Game.CENTURY_BREAK and not self.is100:
            self.show_century_animation()
            self.is100 = True

    def add_foul(self, value):
        """Add foul points to opponent."""
        snap_p1, snap_p2 = self.snapshot_break_stats()

        if self.current_player == 1:
            prev_break = self.break_p1
            self.action_stack.append(("foul", value, self.current_player, prev_break, snap_p1, snap_p2))
            self.score_p2 += value
            self.break_p1 = 0
            self.clear_break_stats_for_player(1)
        else:
            prev_break = self.break_p2
            self.action_stack.append(("foul", value, self.current_player, prev_break, snap_p1, snap_p2))
            self.score_p1 += value
            self.break_p2 = 0
            self.clear_break_stats_for_player(2)

        self.update_scores()
        self.update_remaining_display()

    # ==================== Undo ====================

    def undo_action(self):
        """Undo the last action."""
        if not self.action_stack:
            return

        action = self.action_stack.pop()
        action_type = action[0]

        if action_type == "pot":
            self._undo_pot(action)
        elif action_type == "foul":
            self._undo_foul(action)
        elif action_type == "red_minus":
            prev_red = action[1]
            self.red_remaining = prev_red
            self.ball_remaining_labels[1].setText(str(prev_red))
        elif action_type == "manual":
            self._undo_manual(action)
        elif action_type == "switch":
            self._undo_switch(action)
        elif action_type == "set_is100":
            self.is100 = action[1]

        self.update_scores()
        self.update_remaining_display()

    def _undo_pot(self, action):
        """Undo a pot action."""
        value, player, prev_red, prev_colors, prev_flag, prev_high1, prev_high2, prev_is100 = action[1:]
        self.is100 = prev_is100

        if player == 1:
            self.score_p1 -= value
            if self.current_player == 1:
                self.break_p1 -= value
            label = self.p1_ball_stats.get(value)
        else:
            self.score_p2 -= value
            if self.current_player == 2:
                self.break_p2 -= value
            label = self.p2_ball_stats.get(value)

        if label:
            count = label.text()
            if count and int(count) > 1:
                label.setText(str(int(count) - 1))
            else:
                label.setText("")
                label.hide()

        self.red_remaining = prev_red
        self.colors_remaining = prev_colors
        self.just_potted_last_red = prev_flag
        self.highest_break_p1 = prev_high1
        self.highest_break_p2 = prev_high2
        self.label_high_break_p1.setText(f"Highest: {self.highest_break_p1}")
        self.label_high_break_p2.setText(f"Highest: {self.highest_break_p2}")

        self.decrement_total_ball_stats(player, value)

        if value == 1:
            self.ball_remaining_labels[1].setText(str(self.red_remaining))
        elif value in self.colors_remaining:
            remaining = self.colors_remaining.get(value, 0)
            label_rem = self.ball_remaining_labels.get(value)
            if label_rem:
                label_rem.setText(str(remaining))
                label_rem.setVisible(remaining > 0)

    def _undo_foul(self, action):
        """Undo a foul action."""
        if len(action) >= 6:
            value, player, prev_break, snap_p1, snap_p2 = action[1:]
            self.restore_break_stats_from_snapshot(snap_p1, snap_p2)
        else:
            value, player, prev_break = action[1:]

        if player == 1:
            self.score_p2 -= value
            self.break_p1 = prev_break
        else:
            self.score_p1 -= value
            self.break_p2 = prev_break

    def _undo_manual(self, action):
        """Undo a manual score action."""
        value, player, prev_high1, prev_high2 = action[1:]

        if player == 1:
            self.score_p1 -= value
            if self.current_player == 1:
                self.break_p1 -= value
            label = self.p1_ball_stats.get(value)
        else:
            self.score_p2 -= value
            if self.current_player == 2:
                self.break_p2 -= value
            label = self.p2_ball_stats.get(value)

        if label:
            count = label.text()
            if count and int(count) > 1:
                label.setText(str(int(count) - 1))
            else:
                label.setText("")
                label.hide()

        self.highest_break_p1 = prev_high1
        self.highest_break_p2 = prev_high2
        self.label_high_break_p1.setText(f"Highest: {self.highest_break_p1}")
        self.label_high_break_p2.setText(f"Highest: {self.highest_break_p2}")

        self.decrement_total_ball_stats(player, value)

    def _undo_switch(self, action):
        """Undo a switch player action."""
        if len(action) >= 6:
            prev_player, prev_break_p1, prev_break_p2, snap_p1, snap_p2 = action[1:]
            self.restore_break_stats_from_snapshot(snap_p1, snap_p2)
        else:
            prev_player, prev_break_p1, prev_break_p2 = action[1:]

        self.current_player = prev_player
        self.break_p1 = prev_break_p1
        self.break_p2 = prev_break_p2
        self.update_arrow()

    # ==================== Player Switching ====================

    def switch_player(self):
        """Switch the current player."""
        if self.switch_cooldown:
            return

        self.switch_cooldown = True
        QTimer.singleShot(UI.SWITCH_COOLDOWN, lambda: setattr(self, "switch_cooldown", False))

        prev_player = self.current_player
        prev_break_p1 = self.break_p1
        prev_break_p2 = self.break_p2
        snap_p1, snap_p2 = self.snapshot_break_stats()

        self.action_stack.append(("switch", prev_player, prev_break_p1, prev_break_p2, snap_p1, snap_p2))

        if self.red_remaining == 0:
            self.just_potted_last_red = False

        self.current_player = 2 if self.current_player == 1 else 1

        if self.current_player == 1:
            self.break_p2 = 0
            self.play_switch_arrow_animation("right_to_left")
        else:
            self.break_p1 = 0
            self.play_switch_arrow_animation("left_to_right")

        self.clear_break_stats_for_player(prev_player)
        self.update_scores()

    # ==================== Frame Management ====================

    def reset_scores(self):
        """End current frame and reset for next frame."""
        self.in_black_ball_decider = False
        updated = False
        winning_score = self.total_frames // 2 + 1

        if self.score_p1 > self.score_p2:
            self.frame_score_p1 += 1
            self.frame_scores_p1.append(self.score_p1)
            self.frame_scores_p2.append(self.score_p2)
            updated = True
        elif self.score_p2 > self.score_p1:
            self.frame_score_p2 += 1
            self.frame_scores_p1.append(self.score_p1)
            self.frame_scores_p2.append(self.score_p2)
            updated = True

        if updated:
            self.update_scores()

        self.break_scores_p1.append(self.frame_high_break_p1)
        self.break_scores_p2.append(self.frame_high_break_p2)

        self._reset_frame_state()

        if self.frame_score_p1 >= winning_score:
            QMessageBox.information(self, "比赛结束", f"{self.name_input_p1.text()} 获胜！")
            self.disable_buttons_except_reset()
        elif self.frame_score_p2 >= winning_score:
            QMessageBox.information(self, "比赛结束", f"{self.name_input_p2.text()} 获胜！")
            self.disable_buttons_except_reset()

    def _reset_frame_state(self):
        """Reset state for a new frame."""
        self.frame_high_break_p1 = 0
        self.frame_high_break_p2 = 0
        self.score_p1 = 0
        self.score_p2 = 0
        self.break_p1 = 0
        self.break_p2 = 0
        self.is100 = False
        self.red_remaining = Game.INITIAL_REDS
        self.colors_remaining = {2: 1, 3: 1, 4: 1, 5: 1, 6: 1, 7: 1}

        self.clear_total_stats_both()
        for stats in [self.p1_ball_stats, self.p2_ball_stats]:
            for label in stats.values():
                label.setText("")
                label.hide()

        self.action_stack.clear()
        self.update_scores()
        self.update_remaining_display()

        self.ball_remaining_labels[1].setText(str(self.red_remaining))
        for color_val in range(2, 8):
            remaining = self.colors_remaining[color_val]
            self.ball_remaining_labels[color_val].setText(str(remaining))
            self.ball_remaining_labels[color_val].setVisible(remaining > 0)

        self.match_high_break_start_p1 = self.highest_break_p1
        self.match_high_break_start_p2 = self.highest_break_p2

        self.current_frame_break_player = self._compute_next_break_player()
        self.current_player = self.current_frame_break_player
        self.update_arrow()

    def rerack_frame(self):
        """Re-rack current frame without changing match score."""
        reply = QMessageBox.question(
            self, "确认重摆",
            "确定要本局重摆吗？本局所有得分、单杆、球数统计将被清空，但对局比分保持不变。",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        self.in_black_ball_decider = False
        self.just_potted_last_red = False
        self.is100 = False
        self.score_p1 = 0
        self.score_p2 = 0
        self.break_p1 = 0
        self.break_p2 = 0
        self.frame_high_break_p1 = 0
        self.frame_high_break_p2 = 0
        self.red_remaining = Game.INITIAL_REDS
        self.colors_remaining = {2: 1, 3: 1, 4: 1, 5: 1, 6: 1, 7: 1}

        for stats in [self.p1_ball_stats, self.p2_ball_stats]:
            for label in stats.values():
                label.setText("")
                label.hide()

        self.ball_remaining_labels[1].setText(str(self.red_remaining))
        for color_val in range(2, 8):
            remaining = self.colors_remaining[color_val]
            self.ball_remaining_labels[color_val].setText(str(remaining))
            self.ball_remaining_labels[color_val].setVisible(True)

        self.action_stack.clear()
        self.update_arrow()
        self.update_scores()
        self.update_remaining_display()

        self.highest_break_p1 = self.match_high_break_start_p1
        self.highest_break_p2 = self.match_high_break_start_p2
        self.label_high_break_p1.setText(f"Highest: {self.highest_break_p1}")
        self.label_high_break_p2.setText(f"Highest: {self.highest_break_p2}")

        self.clear_total_stats_both()
        self.current_player = self.current_frame_break_player
        self.update_arrow()

    def full_reset(self):
        """Start a completely new match."""
        self.new_window = SnookerScoreboard()
        self.new_window.show()
        self.close()

    def disable_buttons_except_reset(self):
        """Disable most buttons after match ends."""
        for btn in self.buttons:
            btn.setEnabled(False)
        self.btn_full_reset.setEnabled(True)
        self.btn_stats.setEnabled(True)
        if hasattr(self, "btn_export"):
            self.btn_export.setEnabled(True)
        if hasattr(self, "btn_save"):
            self.btn_save.setEnabled(True)
        if hasattr(self, "btn_load"):
            self.btn_load.setEnabled(True)

    def remove_red_ball(self):
        """Manually remove a red ball."""
        if self.red_remaining > 0:
            self.action_stack.append(("red_minus", self.red_remaining))
            self.red_remaining -= 1
            self.update_remaining_display()
            self.ball_remaining_labels[1].setText(str(self.red_remaining))

    # ==================== Flag Management ====================

    def change_flag(self, player_id):
        """Change player flag."""
        file_name, _ = QFileDialog.getOpenFileName(
            self, "选择国旗图像", ".", "Images (*.svg)"
        )
        if file_name:
            if player_id == 1:
                self.flag_p1.load(file_name)
                self.flag_p1_path = file_name
            elif player_id == 2:
                self.flag_p2.load(file_name)
                self.flag_p2_path = file_name

    @staticmethod
    def load_flag(path, target_size: QSize) -> QPixmap:
        """Load a flag image and scale it."""
        if path.lower().endswith(".svg"):
            image = QImage(target_size, QImage.Format.Format_ARGB32)
            image.fill(Qt.GlobalColor.transparent)
            svg_renderer = QSvgRenderer(path)
            painter = QPainter(image)
            svg_renderer.render(painter)
            painter.end()
            return QPixmap.fromImage(image)
        else:
            image = QImage(path)
            image.setDevicePixelRatio(1.0)
            scaled_image = image.scaled(
                target_size,
                Qt.AspectRatioMode.KeepAspectRatio,
                Qt.TransformationMode.SmoothTransformation,
            )
            return QPixmap.fromImage(scaled_image)

    # ==================== Animations ====================

    def _track_animation(self, anim):
        """Track an animation to prevent premature garbage collection."""
        self._active_animations.append(anim)
        anim.finished.connect(lambda: self._active_animations.remove(anim) if anim in self._active_animations else None)

    def animate_appearance(self, label, from_below: bool = True, distance_px: int | None = None, duration: int = UI.FLY_IN_DURATION):
        """Animate a label appearing with fly-in effect."""
        parent = label.parentWidget()
        if parent is not None:
            lay = parent.layout()
            if lay is not None:
                lay.activate()

        label.show()
        rect = label.geometry()
        if rect.width() <= 0 or rect.height() <= 0:
            rect.setWidth(label.sizeHint().width() or 20)
            rect.setHeight(label.sizeHint().height() or 20)

        dist = distance_px if distance_px is not None else max(10, int(rect.height() * 0.9))
        start_rect = QRect(
            rect.x(),
            rect.y() + (dist if from_below else -dist),
            rect.width(),
            rect.height(),
        )

        eff = QGraphicsOpacityEffect(label)
        label.setGraphicsEffect(eff)
        eff.setOpacity(0.0)

        pos_anim = QPropertyAnimation(label, b"geometry", self)
        pos_anim.setStartValue(start_rect)
        pos_anim.setEndValue(rect)
        pos_anim.setDuration(duration)
        pos_anim.setEasingCurve(getattr(self, "ease", {}).get("fly_pos", QEasingCurve.Type.OutBack))

        fade_anim = QPropertyAnimation(eff, b"opacity", self)
        fade_anim.setStartValue(0.0)
        fade_anim.setEndValue(1.0)
        fade_anim.setDuration(max(120, duration - 40))
        fade_anim.setEasingCurve(getattr(self, "ease", {}).get("fly_opacity", QEasingCurve.Type.OutQuad))

        group = QParallelAnimationGroup(self)
        group.addAnimation(pos_anim)
        group.addAnimation(fade_anim)
        group.finished.connect(lambda: label.setGraphicsEffect(None))

        self._track_animation(group)
        group.start()

    def animate_stat_update(self, label, new_text, scale: float = 1.18):
        """Animate a stat label update with bounce effect."""
        parent = label.parentWidget()
        if parent is not None and parent.layout() is not None:
            parent.layout().activate()

        label.setText(new_text)
        base_rect = label.geometry()
        ss = label.styleSheet()

        def pick(key, default):
            for part in ss.split(";"):
                part = part.strip()
                if part.lower().startswith(key):
                    return part.split(":", 1)[1].strip()
            return default

        def pick_px(key):
            val = pick(key, "")
            digits = "".join(ch for ch in val if ch.isdigit())
            return int(digits) if digits else 0

        bg = QColor(pick("background-color", "#000"))
        fg = QColor(pick("color", "#fff"))
        pad_top = pick_px("padding-top")

        ghost = BadgeGhost(base_rect, bg, fg, new_text, label.font(), parent, max_scale=scale, text_pad_top=pad_top)
        label.setVisible(False)

        grow = QPropertyAnimation(ghost, b"factor", self)
        grow.setDuration(UI.STAT_UPDATE_DURATION)
        grow.setStartValue(1.0)
        grow.setEndValue(scale)
        grow.setEasingCurve(QEasingCurve.Type.OutCubic)

        pause = QPauseAnimation(100)

        shrink = QPropertyAnimation(ghost, b"factor", self)
        shrink.setDuration(180)
        shrink.setStartValue(scale)
        shrink.setEndValue(1.0)
        shrink.setEasingCurve(QEasingCurve.Type.InCubic)

        seq = QSequentialAnimationGroup(self)
        seq.addAnimation(grow)
        seq.addAnimation(pause)
        seq.addAnimation(shrink)

        def finish():
            ghost.deleteLater()
            label.setVisible(True)

        seq.finished.connect(finish)
        self._track_animation(seq)
        seq.start()

    def play_switch_arrow_animation(self, direction: str):
        """Play arrow switching animation."""
        if direction == "left_to_right":
            source_arrow = self.arrow_left
            target_arrow = self.arrow_right
        else:
            source_arrow = self.arrow_right
            target_arrow = self.arrow_left

        opacity_effect = QGraphicsOpacityEffect(target_arrow)
        opacity_effect.setOpacity(0.0)
        target_arrow.setGraphicsEffect(opacity_effect)

        source_start = source_arrow.pos()
        offset = 60 if direction == "left_to_right" else -60
        source_end = QPoint(source_start.x() + offset, source_start.y())

        anim_out = QPropertyAnimation(source_arrow, b"pos")
        anim_out.setDuration(UI.ARROW_ANIMATION_DURATION)
        anim_out.setStartValue(source_start)
        anim_out.setEndValue(source_end)
        anim_out.setEasingCurve(QEasingCurve.Type.InBack)

        if direction == "right_to_left":
            target_end = target_arrow.pos()
        else:
            target_end = QPoint(643, 0)

        target_offset = -60 if direction == "left_to_right" else 60
        target_start = QPoint(target_end.x() + target_offset, target_end.y())

        self.label_score_p2.raise_()
        target_arrow.move(target_start)
        target_arrow.show()

        anim_in = QPropertyAnimation(target_arrow, b"pos")
        anim_in.setDuration(UI.ARROW_ANIMATION_DURATION)
        anim_in.setStartValue(target_start)
        anim_in.setEndValue(target_end)
        anim_in.setEasingCurve(QEasingCurve.Type.OutBack)

        fade_animation = QPropertyAnimation(opacity_effect, b"opacity", self)
        fade_animation.setStartValue(0.0)
        fade_animation.setEndValue(1.0)
        fade_animation.setDuration(1)

        group1 = QParallelAnimationGroup(self)
        group1.addAnimation(anim_in)
        group1.addAnimation(fade_animation)

        group = QSequentialAnimationGroup(self)
        group.addAnimation(anim_out)
        group.addAnimation(QPauseAnimation(25))
        group.addAnimation(group1)

        def after_anim():
            source_arrow.hide()
            source_arrow.move(source_start)

        group.finished.connect(after_anim)
        self._track_animation(group)
        group.start()

    def show_century_animation(self):
        """Show century celebration animation."""
        self.century_label.show()
        self.century_label.raise_()

        bar_width = self.break_bar.width()
        bar_height = self.break_bar.height()
        label_width = 200
        label_height = 50
        x_center = (bar_width - label_width) // 2
        y_center = (bar_height - label_height) // 2

        grow_anim = QPropertyAnimation(self.century_label, b"geometry")
        grow_anim.setDuration(UI.CENTURY_GROW_DURATION)
        grow_anim.setStartValue(QRect(x_center + label_width // 2, y_center + label_height // 2, 0, 0))
        grow_anim.setEndValue(QRect(x_center, y_center, label_width, label_height))
        grow_anim.setEasingCurve(QEasingCurve.Type.OutCirc)

        pause = QPauseAnimation(UI.CENTURY_PAUSE_DURATION)

        shrink_anim = QPropertyAnimation(self.century_label, b"geometry")
        shrink_anim.setDuration(UI.CENTURY_SHRINK_DURATION)
        shrink_anim.setStartValue(QRect(x_center, y_center, label_width, label_height))
        shrink_anim.setEndValue(QRect(x_center + label_width // 2, y_center + label_height // 2, 0, 0))
        shrink_anim.setEasingCurve(QEasingCurve.Type.InBack)
        shrink_anim.finished.connect(self.century_label.hide)

        anim_group = QSequentialAnimationGroup(self)
        anim_group.addAnimation(grow_anim)
        anim_group.addAnimation(pause)
        anim_group.addAnimation(shrink_anim)

        self.century_animation = anim_group
        self._track_animation(anim_group)
        anim_group.start()

    # ==================== Button Handling ====================

    def handle_score_click(self, value, button):
        """Handle ball button click with cooldown."""
        if self.score_cooldowns.get(value, False):
            return
        self.pot_ball(value)
        self.score_cooldowns[value] = True
        button.setEnabled(False)
        QTimer.singleShot(UI.BALL_COOLDOWN, lambda: self.reset_cooldown(value, button))

    def reset_cooldown(self, value, button):
        """Reset button cooldown."""
        self.score_cooldowns[value] = False
        button.setEnabled(True)

    def clear_score_cooldown(self, value):
        """Clear score cooldown."""
        self.score_cooldowns[value] = False

    # ==================== Keyboard Events ====================

    def keyPressEvent(self, event):
        """Handle keyboard shortcuts."""
        key = event.key()
        modifiers = event.modifiers()

        if modifiers & Qt.KeyboardModifier.ControlModifier and key == Qt.Key.Key_Z:
            self.undo_action()
            return

        if key == Qt.Key.Key_Space:
            self.switch_player()
            return

        if modifiers & Qt.KeyboardModifier.ControlModifier and key in (Qt.Key.Key_4, Qt.Key.Key_5, Qt.Key.Key_6, Qt.Key.Key_7):
            self.add_foul(key - Qt.Key.Key_0)
            return

        if key in (Qt.Key.Key_1, Qt.Key.Key_2, Qt.Key.Key_3, Qt.Key.Key_4, Qt.Key.Key_5, Qt.Key.Key_6, Qt.Key.Key_7):
            self.pot_ball(int(key) - Qt.Key.Key_0)
            return

        if modifiers & Qt.KeyboardModifier.ControlModifier and key == Qt.Key.Key_R:
            self.full_reset()
            return

        if key == Qt.Key.Key_R:
            self.reset_scores()
            return

        if key == Qt.Key.Key_Minus:
            self.remove_red_ball()
            return

        if key == Qt.Key.Key_H:
            self.toggle_overlay()
            return

    # ==================== Statistics ====================

    def show_statistics(self):
        """Show match statistics dialog."""
        dialog = QDialog(self)
        dialog.setWindowTitle("每局得分统计")
        layout = QVBoxLayout(dialog)

        col_count = max(len(self.frame_scores_p1), len(self.frame_scores_p2))
        table = QTableWidget()
        table.setRowCount(2)
        table.setColumnCount(col_count)
        table.setVerticalHeaderLabels([self.name_input_p1.text(), self.name_input_p2.text()])

        for i in range(col_count):
            score1 = self.frame_scores_p1[i] if i < len(self.frame_scores_p1) else 0
            score2 = self.frame_scores_p2[i] if i < len(self.frame_scores_p2) else 0
            break1 = self.break_scores_p1[i] if i < len(self.break_scores_p1) else 0
            break2 = self.break_scores_p2[i] if i < len(self.break_scores_p2) else 0

            text1 = f"{score1} (c)" if break1 >= Game.CENTURY_BREAK else str(score1)
            text2 = f"{score2} (c)" if break2 >= Game.CENTURY_BREAK else str(score2)

            item1 = QTableWidgetItem(text1)
            item2 = QTableWidgetItem(text2)

            if score1 > score2:
                item1.setBackground(QColor(Styles.WINNER_HIGHLIGHT))
            elif score2 > score1:
                item2.setBackground(QColor(Styles.WINNER_HIGHLIGHT))

            table.setItem(0, i, item1)
            table.setItem(1, i, item2)

        layout.addWidget(table)
        dialog.setLayout(layout)
        dialog.resize(400, 200)
        dialog.exec()

    # ==================== Export ====================

    def export_to_excel(self):
        """Export match data to Excel."""
        p1 = (self.name_input_p1.text() or "Player 1").strip()
        p2 = (self.name_input_p2.text() or "Player 2").strip()
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        suggested = f"{p1}_vs_{p2}_{ts}.xlsx"

        path, _ = QFileDialog.getSaveFileName(self, "导出为 Excel", suggested, "Excel 文件 (*.xlsx)")
        if not path:
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            QMessageBox.warning(self, "缺少依赖", "需要安装 openpyxl 才能导出 Excel：\n\npip install openpyxl")
            return

        finished_frames = self._collect_frame_data(p1, p2)

        wb = Workbook()
        ws = wb.active
        ws.title = "总览"

        self._write_overview_sheet(ws, p1, p2, finished_frames, Font, Alignment, PatternFill, get_column_letter)

        ws2 = wb.create_sheet("分局明细")
        self._write_detail_sheet(ws2, p1, p2, finished_frames, Font, Alignment, PatternFill, get_column_letter)

        try:
            wb.save(path)
        except Exception as e:
            QMessageBox.critical(self, "导出失败", f"写入文件时出错：\n{e}")
            return

        QMessageBox.information(self, "导出完成", f"已导出到：\n{path}")

    def _collect_frame_data(self, p1, p2):
        """Collect frame data for export."""
        finished_frames = []
        n = max(len(self.frame_scores_p1), len(self.frame_scores_p2))

        for i in range(n):
            s1 = self.frame_scores_p1[i] if i < len(self.frame_scores_p1) else 0
            s2 = self.frame_scores_p2[i] if i < len(self.frame_scores_p2) else 0
            b1 = self.break_scores_p1[i] if i < len(self.break_scores_p1) else 0
            b2 = self.break_scores_p2[i] if i < len(self.break_scores_p2) else 0
            winner = p1 if s1 > s2 else (p2 if s2 > s1 else "平局")

            finished_frames.append({
                "Frame": i + 1,
                f"{p1} 分数": s1,
                f"{p2} 分数": s2,
                "胜者": winner,
                f"{p1} 单杆最高": b1,
                f"{p2} 单杆最高": b2,
                f"{p1} Century?": "✔" if b1 >= Game.CENTURY_BREAK else "",
                f"{p2} Century?": "✔" if b2 >= Game.CENTURY_BREAK else "",
                "状态": "已结束",
            })

        # Add current frame if in progress
        has_current = (
            self.score_p1 > 0 or self.score_p2 > 0 or
            self.red_remaining != Game.INITIAL_REDS or
            any(self.colors_remaining.get(k, 0) != 1 for k in (2, 3, 4, 5, 6, 7)) or
            any((lbl.text() or "") for lbl in self.p1_ball_stats.values()) or
            any((lbl.text() or "") for lbl in self.p2_ball_stats.values())
        )

        if has_current:
            finished_count = len(finished_frames)
            finished_frames.append({
                "Frame": finished_count + 1,
                f"{p1} 分数": self.score_p1,
                f"{p2} 分数": self.score_p2,
                "胜者": "进行中",
                f"{p1} 单杆最高": self.frame_high_break_p1,
                f"{p2} 单杆最高": self.frame_high_break_p2,
                f"{p1} Century?": "✔" if self.frame_high_break_p1 >= Game.CENTURY_BREAK else "",
                f"{p2} Century?": "✔" if self.frame_high_break_p2 >= Game.CENTURY_BREAK else "",
                "状态": "进行中",
            })

        return finished_frames

    def _write_overview_sheet(self, ws, p1, p2, finished_frames, Font, Alignment, PatternFill, get_column_letter):
        """Write overview sheet to Excel."""
        ws.merge_cells("A1:F1")
        ws["A1"] = f"斯诺克比赛统计：{p1} vs {p2}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws["A3"], ws["B3"] = "导出时间", datetime.now().strftime("%Y-%m-%d %H:%M")
        ws["A4"], ws["B4"] = "Best of", self.total_frames
        ws["A5"], ws["B5"] = "对局比分", f"{self.frame_score_p1} - {self.frame_score_p2}"
        ws["A6"], ws["B6"] = p1 + " Highest", self.highest_break_p1
        ws["A7"], ws["B7"] = p2 + " Highest", self.highest_break_p2

        winning_score = self.total_frames // 2 + 1
        if self.frame_score_p1 >= winning_score:
            status = p1
        elif self.frame_score_p2 >= winning_score:
            status = p2
        else:
            status = "进行中"
        ws["D3"], ws["E3"] = "比赛状态", status

        ws["A9"] = "每局概览"
        ws["A9"].font = Font(bold=True)

        overview_headers = ["Frame", f"{p1} 分数", f"{p2} 分数", "胜者/状态"]
        for c, h in enumerate(overview_headers, start=1):
            cell = ws.cell(row=10, column=c)
            cell.value = h
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill("solid", fgColor="FFD110")

        r = 11
        for row in finished_frames:
            ws.cell(row=r, column=1, value=row["Frame"]).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=row[f"{p1} 分数"]).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3, value=row[f"{p2} 分数"]).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=4, value=row.get("胜者", row.get("状态", ""))).alignment = Alignment(horizontal="center")
            r += 1

        for col in range(1, len(overview_headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 14

    def _write_detail_sheet(self, ws2, p1, p2, finished_frames, Font, Alignment, PatternFill, get_column_letter):
        """Write detail sheet to Excel."""
        headers = [
            "Frame", f"{p1} 分数", f"{p2} 分数", "胜者",
            f"{p1} 单杆最高", f"{p2} 单杆最高",
            f"{p1} Century?", f"{p2} Century?", "状态",
        ]
        ws2.append(headers)

        head_fill = PatternFill("solid", fgColor="FFD110")
        for c in range(1, len(headers) + 1):
            cell = ws2.cell(row=1, column=c)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = head_fill

        winner_fill = PatternFill("solid", fgColor="FFF275")
        for row in finished_frames:
            s1 = row[f"{p1} 分数"]
            s2 = row[f"{p2} 分数"]
            b1 = row[f"{p1} 单杆最高"]
            b2 = row[f"{p2} 单杆最高"]

            s1_text = f"{s1} (c)" if (isinstance(b1, (int, float)) and b1 >= Game.CENTURY_BREAK) else str(s1)
            s2_text = f"{s2} (c)" if (isinstance(b2, (int, float)) and b2 >= Game.CENTURY_BREAK) else str(s2)

            r = ws2.max_row + 1
            ws2.cell(row=r, column=1, value=row["Frame"])
            ws2.cell(row=r, column=2, value=s1_text)
            ws2.cell(row=r, column=3, value=s2_text)
            ws2.cell(row=r, column=4, value=row["胜者"])
            ws2.cell(row=r, column=5, value=b1)
            ws2.cell(row=r, column=6, value=b2)
            ws2.cell(row=r, column=7, value=row.get(f"{p1} Century?", ""))
            ws2.cell(row=r, column=8, value=row.get(f"{p2} Century?", ""))
            ws2.cell(row=r, column=9, value=row["状态"])

            for c in range(1, 10):
                ws2.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center")

            if row["状态"] == "已结束":
                if row["胜者"] == p1:
                    ws2.cell(row=r, column=2).fill = winner_fill
                elif row["胜者"] == p2:
                    ws2.cell(row=r, column=3).fill = winner_fill

        for col in range(1, len(headers) + 1):
            letter = get_column_letter(col)
            max_len = max(len(str(ws2.cell(row=r, column=col).value or "")) for r in range(1, ws2.max_row + 1))
            ws2.column_dimensions[letter].width = max(10, min(28, max_len + 2))

    # ==================== Save/Load ====================

    def _collect_counts(self, stats_dict: dict) -> dict:
        """Collect counts from stats labels."""
        out = {}
        for v, lbl in stats_dict.items():
            t = lbl.text()
            out[str(v)] = int(t) if t else 0
        return out

    def _apply_counts(self, stats_dict: dict, counts: dict):
        """Apply counts to stats labels."""
        for v, lbl in stats_dict.items():
            n = int(counts.get(str(v), 0)) if counts else 0
            if n > 0:
                lbl.setText(str(n))
                lbl.show()
            else:
                lbl.setText("")
                lbl.hide()

    def save_game(self):
        """Save game state to file."""
        p1 = (self.name_input_p1.text() or "Player1").strip()
        p2 = (self.name_input_p2.text() or "Player2").strip()
        ts = datetime.now().strftime("%Y%m%d_%H%M")
        suggested = f"{p1}_vs_{p2}_{ts}.snooker.json"

        path, _ = QFileDialog.getSaveFileName(self, "保存存档", suggested, "Snooker 存档 (*.json *.snooker.json)")
        if not path:
            return

        colors = {str(k): int(v) for k, v in self.colors_remaining.items()}
        break_counts_p1 = self._collect_counts(self.p1_ball_stats)
        break_counts_p2 = self._collect_counts(self.p2_ball_stats)
        frame_total_counts_p1 = self._collect_counts(getattr(self, "p1_frame_total_stats", {}))
        frame_total_counts_p2 = self._collect_counts(getattr(self, "p2_frame_total_stats", {}))

        state = {
            "meta": {"version": 1, "saved_at": ts},
            "names": [self.name_input_p1.text(), self.name_input_p2.text()],
            "flags": [
                getattr(self, "flag_p1_path", get_resource_path(Resources.FLAG_P1_DEFAULT)),
                getattr(self, "flag_p2_path", get_resource_path(Resources.FLAG_P2_DEFAULT)),
            ],
            "overlay_pos": [
                self.overlay.x() if hasattr(self, "overlay") else 100,
                self.overlay.y() if hasattr(self, "overlay") else 100,
            ],
            "match": {
                "total_frames": self.total_frames,
                "frame_score_p1": self.frame_score_p1,
                "frame_score_p2": self.frame_score_p2,
                "frame_scores_p1": self.frame_scores_p1,
                "frame_scores_p2": self.frame_scores_p2,
                "break_scores_p1": self.break_scores_p1,
                "break_scores_p2": self.break_scores_p2,
                "highest_break_p1": self.highest_break_p1,
                "highest_break_p2": self.highest_break_p2,
                "match_high_break_start_p1": getattr(self, "match_high_break_start_p1", self.highest_break_p1),
                "match_high_break_start_p2": getattr(self, "match_high_break_start_p2", self.highest_break_p2),
                "first_break_player": getattr(self, "first_break_player", 1),
                "current_frame_break_player": getattr(self, "current_frame_break_player", getattr(self, "first_break_player", 1)),
            },
            "frame": {
                "current_player": self.current_player,
                "score_p1": self.score_p1,
                "score_p2": self.score_p2,
                "break_p1": self.break_p1,
                "break_p2": self.break_p2,
                "frame_high_break_p1": self.frame_high_break_p1,
                "frame_high_break_p2": self.frame_high_break_p2,
                "red_remaining": self.red_remaining,
                "colors_remaining": colors,
                "just_potted_last_red": self.just_potted_last_red,
                "in_black_ball_decider": self.in_black_ball_decider,
                "is100": self.is100,
                "break_counts_p1": break_counts_p1,
                "break_counts_p2": break_counts_p2,
                "frame_total_counts_p1": frame_total_counts_p1,
                "frame_total_counts_p2": frame_total_counts_p2,
            },
        }

        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(state, f, ensure_ascii=False, indent=2)
        except Exception as e:
            QMessageBox.critical(self, "保存失败", f"写入文件时出错：\n{e}")
            return

        QMessageBox.information(self, "已保存", f"存档成功：\n{path}")

    def _as_int(self, v, default=0) -> int:
        """Convert value to int with fallback."""
        try:
            return int(v)
        except Exception:
            try:
                return int(float(v))
            except Exception:
                return default

    def capture_arrow_positions(self):
        """Capture initial arrow positions."""
        if not self.arrow_left.isVisible() or not self.arrow_right.isVisible():
            QTimer.singleShot(0, self.capture_arrow_positions)
            return
        self.arrow_left_initial_pos = self.arrow_left.pos()
        self.arrow_right_initial_pos = self.arrow_right.pos()

    def load_game(self):
        """Load game state from file."""
        path, _ = QFileDialog.getOpenFileName(self, "加载存档", ".", "Snooker 存档 (*.json *.snooker.json)")
        if not path:
            return

        try:
            with open(path, "r", encoding="utf-8") as f:
                state = json.load(f)
        except Exception as e:
            QMessageBox.critical(self, "加载失败", f"读取存档出错：\n{e}")
            return

        try:
            self._apply_loaded_state(state)
            QTimer.singleShot(0, lambda: getattr(self, "capture_arrow_positions", lambda: None)())
            QMessageBox.information(self, "加载完成", f"已恢复存档：\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "加载失败", f"解析存档时出错：\n{e}")

    def _apply_loaded_state(self, state):
        """Apply loaded state to the scoreboard."""
        match = state.get("match", {}) or {}
        frame = state.get("frame", {}) or {}
        names = state.get("names", ["Player 1", "Player 2"])

        self.name_input_p1.setText(str(names[0]) if len(names) > 0 else "Player 1")
        self.name_input_p2.setText(str(names[1]) if len(names) > 1 else "Player 2")

        flags = state.get("flags", [get_resource_path(Resources.FLAG_P1_DEFAULT), get_resource_path(Resources.FLAG_P2_DEFAULT)])
        f1 = flags[0] if len(flags) > 0 else get_resource_path(Resources.FLAG_P1_DEFAULT)
        f2 = flags[1] if len(flags) > 1 else get_resource_path(Resources.FLAG_P2_DEFAULT)

        if isinstance(f1, str) and os.path.exists(f1):
            self.flag_p1.load(f1)
            self.flag_p1_path = f1
        if isinstance(f2, str) and os.path.exists(f2):
            self.flag_p2.load(f2)
            self.flag_p2_path = f2

        self.total_frames = self._as_int(match.get("total_frames", self.total_frames), self.total_frames)
        self.label_frame_center.setText(f"({self.total_frames})")

        self.frame_score_p1 = self._as_int(match.get("frame_score_p1", 0), 0)
        self.frame_score_p2 = self._as_int(match.get("frame_score_p2", 0), 0)
        self.frame_scores_p1 = [self._as_int(x, 0) for x in match.get("frame_scores_p1", [])]
        self.frame_scores_p2 = [self._as_int(x, 0) for x in match.get("frame_scores_p2", [])]
        self.break_scores_p1 = [self._as_int(x, 0) for x in match.get("break_scores_p1", [])]
        self.break_scores_p2 = [self._as_int(x, 0) for x in match.get("break_scores_p2", [])]

        self.highest_break_p1 = self._as_int(match.get("highest_break_p1", 0), 0)
        self.highest_break_p2 = self._as_int(match.get("highest_break_p2", 0), 0)
        self.match_high_break_start_p1 = self._as_int(match.get("match_high_break_start_p1", self.highest_break_p1), self.highest_break_p1)
        self.match_high_break_start_p2 = self._as_int(match.get("match_high_break_start_p2", self.highest_break_p2), self.highest_break_p2)

        self.current_player = self._as_int(frame.get("current_player", 1), 1)
        self.score_p1 = self._as_int(frame.get("score_p1", 0), 0)
        self.score_p2 = self._as_int(frame.get("score_p2", 0), 0)
        self.break_p1 = self._as_int(frame.get("break_p1", 0), 0)
        self.break_p2 = self._as_int(frame.get("break_p2", 0), 0)
        self.frame_high_break_p1 = self._as_int(frame.get("frame_high_break_p1", 0), 0)
        self.frame_high_break_p2 = self._as_int(frame.get("frame_high_break_p2", 0), 0)
        self.red_remaining = self._as_int(frame.get("red_remaining", Game.INITIAL_REDS), Game.INITIAL_REDS)

        self.first_break_player = self._as_int(match.get("first_break_player", 1), 1)
        if "current_frame_break_player" in match:
            self.current_frame_break_player = self._as_int(match.get("current_frame_break_player", self.first_break_player), self.first_break_player)
        else:
            self.current_frame_break_player = self._compute_next_break_player()

        colors_str = frame.get("colors_remaining", {}) or {}
        self.colors_remaining = {}
        for k, v in colors_str.items():
            kk = self._as_int(k, None)
            vv = self._as_int(v, 0)
            if kk in (2, 3, 4, 5, 6, 7):
                self.colors_remaining[kk] = vv
        for kk in (2, 3, 4, 5, 6, 7):
            self.colors_remaining.setdefault(kk, 0)

        self.just_potted_last_red = bool(frame.get("just_potted_last_red", False))
        self.in_black_ball_decider = bool(frame.get("in_black_ball_decider", False))
        self.is100 = bool(frame.get("is100", False))

        self._apply_counts(self.p1_ball_stats, frame.get("break_counts_p1", {}) or {})
        self._apply_counts(self.p2_ball_stats, frame.get("break_counts_p2", {}) or {})

        if hasattr(self, "p1_frame_total_stats"):
            self._apply_counts(self.p1_frame_total_stats, frame.get("frame_total_counts_p1", {}) or {})
        if hasattr(self, "p2_frame_total_stats"):
            self._apply_counts(self.p2_frame_total_stats, frame.get("frame_total_counts_p2", {}) or {})

        if 1 in self.ball_remaining_labels:
            self.ball_remaining_labels[1].setText(str(self.red_remaining))
        for color_val in range(2, 8):
            lbl = self.ball_remaining_labels.get(color_val)
            if lbl:
                remaining = self.colors_remaining.get(color_val, 0)
                lbl.setText(str(remaining))
                lbl.setVisible(remaining > 0)

        self.update_arrow()
        self.update_scores()
        self.update_remaining_display()
        self.action_stack.clear()

        if hasattr(self, "overlay"):
            pos = state.get("overlay_pos", [self.overlay.x(), self.overlay.y()]) or [self.overlay.x(), self.overlay.y()]
            try:
                x = self._as_int(pos[0], self.overlay.x())
                y = self._as_int(pos[1], self.overlay.y())
                self.overlay.move(max(0, x), max(0, y))
            except Exception:
                pass

    # ==================== Overlay ====================

    def toggle_overlay(self):
        """Toggle overlay visibility."""
        if self.overlay.isVisible():
            self.animate_hide_overlay()
        else:
            self.animate_show_overlay()

    def animate_hide_overlay(self):
        """Animate hiding the overlay."""
        self.current_x = self.overlay.x()
        self.current_y = self.overlay.y()

        pos_anim = QPropertyAnimation(self.overlay, b"geometry", self)
        pos_anim.setDuration(600)
        pos_anim.setStartValue(QRect(self.current_x, self.current_y, UI.OVERLAY_WIDTH, 200))
        pos_anim.setEndValue(QRect(self.current_x, self.current_y + 40, UI.OVERLAY_WIDTH, 200))
        pos_anim.setEasingCurve(QEasingCurve.Type.InQuint)

        fade_anim = QPropertyAnimation(self.overlay, b"windowOpacity", self)
        fade_anim.setDuration(600)
        fade_anim.setStartValue(1.0)
        fade_anim.setEndValue(0.0)
        fade_anim.setEasingCurve(QEasingCurve.Type.InCirc)

        group = QParallelAnimationGroup(self)
        group.addAnimation(pos_anim)
        group.addAnimation(fade_anim)
        group.finished.connect(self.overlay.hide)

        self._track_animation(group)
        group.start()
        self.overlay_hide_anim = group

    def animate_show_overlay(self):
        """Animate showing the overlay."""
        target_x = getattr(self, "current_x", self.overlay.x())
        target_y = getattr(self, "current_y", self.overlay.y())
        start_y = target_y + 50

        self.overlay.move(target_x, start_y)
        self.overlay.resize(UI.OVERLAY_WIDTH, 200)
        self.overlay.show()

        pos_anim = QPropertyAnimation(self.overlay, b"geometry", self)
        pos_anim.setDuration(1000)
        pos_anim.setStartValue(QRect(target_x, start_y, UI.OVERLAY_WIDTH, 200))
        pos_anim.setEndValue(QRect(target_x, target_y, UI.OVERLAY_WIDTH, 200))
        pos_anim.setEasingCurve(QEasingCurve.Type.OutQuart)

        fade_anim = QPropertyAnimation(self.overlay, b"windowOpacity", self)
        fade_anim.setDuration(1000)
        fade_anim.setStartValue(0.0)
        fade_anim.setEndValue(1.0)
        fade_anim.setEasingCurve(QEasingCurve.Type.OutQuart)

        group = QParallelAnimationGroup(self)
        group.addAnimation(pos_anim)
        group.addAnimation(fade_anim)

        def after_show():
            self.overlay.activateWindow()
            self.overlay.setFocus()
            self.name_input_p1.clearFocus()
            self.name_input_p2.clearFocus()

        group.finished.connect(after_show)
        self._track_animation(group)
        group.start()
        self.overlay_show_anim = group

    # ==================== Layout Helpers ====================

    def _layout_index_of(self, layout, widget) -> int:
        """Find widget index in layout."""
        for i in range(layout.count()):
            it = layout.itemAt(i)
            if it and it.widget() is widget:
                return i
        return -1

    def _push_neighbors_then_fly_in(self, layout, label, align_right: bool):
        """Push neighboring badges then animate fly-in."""
        idx = self._layout_index_of(layout, label)
        if idx < 0:
            self.animate_appearance(label, from_below=False)
            return

        insert_at = idx + 1 if align_right else idx
        h = label.height() or label.sizeHint().height() or 28
        spacer = QSpacerItem(0, h, QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        layout.insertSpacerItem(insert_at, spacer)

        badge_w = label.width() or label.sizeHint().width() or 28
        gap = layout.spacing()
        extra = getattr(self, "break_push_extra", 0)
        target_w = badge_w + gap + extra

        anim_obj = HSpacerAnimator(layout, spacer, h, parent=self)
        anim = QPropertyAnimation(anim_obj, b"width", self)

        if not self._need_push(layout, label, align_right):
            anim.setDuration(1)
        else:
            anim.setDuration(250)

        anim.setStartValue(0)
        anim.setEndValue(target_w)
        anim.setEasingCurve(QEasingCurve.Type.OutCirc)

        def _finish():
            for i in range(layout.count()):
                it = layout.itemAt(i)
                if it and it.spacerItem() is spacer:
                    layout.takeAt(i)
                    break
            self.animate_appearance(label, from_below=False)

        anim.finished.connect(_finish)
        self._track_animation(anim)
        anim.start()

    def _need_push(self, layout, label, align_right: bool) -> bool:
        """Check if push animation is needed."""
        idx = self._layout_index_of(layout, label)
        if idx < 0:
            return False

        if align_right:
            for j in range(0, idx):
                it = layout.itemAt(j)
                w = it.widget() if it else None
                if w is not None and w.isVisible():
                    return True
            return False
        else:
            for j in range(idx + 1, layout.count()):
                it = layout.itemAt(j)
                w = it.widget() if it else None
                if w is not None and w.isVisible():
                    return True
            return False

    # ==================== Global Hotkey ====================

    def _install_global_hotkey(self):
        """Install global hotkey for overlay toggle."""
        if not sys.platform.startswith("win"):
            return

        vk = ord("H")
        mod = GlobalHotkey.MOD_CONTROL | GlobalHotkey.MOD_ALT
        try:
            self._hotkey = GlobalHotkey(self.overlay, vk=vk, modifiers=mod, callback=self.toggle_overlay)
        except Exception as e:
            print("注册全局热键失败：", e)

    # ==================== Window Events ====================

    def closeEvent(self, event):
        """Handle window close event."""
        # Unregister hotkey
        if hasattr(self, "_hotkey"):
            try:
                self._hotkey.unregister()
            except Exception:
                pass

        # Close overlay
        if hasattr(self, "overlay") and self.overlay is not None:
            self.overlay.close()

        super().closeEvent(event)
